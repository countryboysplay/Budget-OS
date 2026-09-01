"""
Budget-OS publisher

Watches the five sales comparison workbooks in Dropbox. When one is saved,
reads the ANALYSIS FROM GR. SLS. REPORT tab, writes data/<location>.json
in the repo, commits, and pushes. Cloudflare Pages redeploys on push.

Usage:
    python publish.py            publish everything once, then watch for changes
    python publish.py --once     publish everything once and exit
    python publish.py --no-push  write JSON and commit, but don't push (for testing)
"""

import argparse
import json
import logging
import os
import queue
import subprocess
import sys
import threading
import time
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import range_boundaries

HERE = Path(__file__).resolve().parent
CONFIG_PATH = HERE / "config.json"
LOG_PATH = HERE / "publisher.log"

# Sheet column letter -> JSON key. Column J is a spacer and is skipped.
COLUMNS = {
    "B": "quotes_booked",
    "C": "quotes_completed",
    "D": "ytd_gross",
    "E": "ytd_net",
    "F": "est_remaining_gross",
    "G": "est_ye_net",
    "H": "ye_budget",
    "I": "pct_of_budget",
    "K": "ytd_net_dollars",
    "L": "est_ye_net_dollars",
    "M": "ye_budget_dollars",
    "N": "est_variance_dollars",
    "O": "est_variance_pct",
}
HEADER_ROWS = 4  # rows 1-4 are the stacked header block

log = logging.getLogger("publisher")


# ---------------------------------------------------------------- helpers

def now_iso() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def clean_text(v):
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def clean_number(v):
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, str):
        s = v.strip().replace(",", "")
        if not s:
            return None
        try:
            v = float(s)
        except ValueError:
            return None
    if isinstance(v, int):
        return v
    if isinstance(v, float):
        r = round(v, 2)
        return int(r) if r.is_integer() else r
    return None


def norm_path(p: str) -> str:
    return os.path.normcase(os.path.abspath(p))


# ---------------------------------------------------------------- reading

def read_sheet(path: str, sheet: str, cell_range: str, retries: int = 5):
    """Return the range as a list of row tuples. Retries while Excel/Dropbox
    still hold the file."""
    min_col, min_row, max_col, max_row = range_boundaries(cell_range)
    last_err = None
    for attempt in range(1, retries + 1):
        try:
            wb = load_workbook(path, read_only=True, data_only=True)
            try:
                if sheet not in wb.sheetnames:
                    raise KeyError(
                        f"Sheet {sheet!r} not found. Sheets: {wb.sheetnames}")
                ws = wb[sheet]
                return [
                    tuple(r)
                    for r in ws.iter_rows(min_row=min_row, max_row=max_row,
                                          min_col=min_col, max_col=max_col,
                                          values_only=True)
                ]
            finally:
                wb.close()
        except (PermissionError, OSError) as e:
            last_err = e
            log.warning("Read attempt %d/%d failed for %s: %s",
                        attempt, retries, Path(path).name, e)
            time.sleep(2 * attempt)
    raise RuntimeError(f"Could not read {path}: {last_err}")


def rows_to_records(rows):
    """Turn the raw grid into tagged records. Sections are detected by
    label so a shifted row number won't break anything."""
    col_index = {letter: ord(letter) - ord("A") for letter in COLUMNS}
    records = []
    section = "programs"

    for row in rows[HEADER_ROWS:]:
        category = clean_text(row[0])
        if category is None:
            continue
        upper = category.upper()

        rec = {"category": category, "section": section, "is_total": False}
        for letter, key in COLUMNS.items():
            rec[key] = clean_number(row[col_index[letter]])

        if upper == "TOTAL":
            rec["section"] = "all"
            rec["is_total"] = True
        elif upper.startswith("TOTAL"):
            rec["is_total"] = True
        elif upper == "INDUSTRIALS":
            rec["section"] = "industrials"

        records.append(rec)

        # Advance the section after the marker rows.
        if upper.startswith("TOTAL PROGRAMS"):
            section = "industrials"
        elif upper == "INDUSTRIALS":
            section = "addons"

    return records


def build_payload(loc: dict, cfg: dict) -> dict:
    sheet = loc.get("sheet", cfg["sheet"])
    cell_range = loc.get("range", cfg["range"])
    rows = read_sheet(loc["path"], sheet, cell_range)
    mtime = datetime.fromtimestamp(os.path.getmtime(loc["path"]),
                                   tz=timezone.utc)
    return {
        "key": loc["key"],
        "label": loc["label"],
        "generated_at": now_iso(),
        "source_modified": mtime.isoformat(timespec="seconds"),
        "rows": rows_to_records(rows),
    }


# ---------------------------------------------------------------- writing

def comparable(payload: dict) -> str:
    """JSON without the timestamps, for change detection."""
    p = {k: v for k, v in payload.items()
         if k not in ("generated_at", "source_modified")}
    return json.dumps(p, sort_keys=True)


def write_if_changed(path: Path, payload: dict) -> bool:
    if path.exists():
        try:
            old = json.loads(path.read_text(encoding="utf-8"))
            if comparable(old) == comparable(payload):
                return False
        except (ValueError, OSError):
            pass
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, indent=1), encoding="utf-8")
    return True


def write_index(data_dir: Path, cfg: dict) -> None:
    entries = []
    for loc in cfg["locations"]:
        f = data_dir / f"{loc['key']}.json"
        updated = None
        if f.exists():
            try:
                updated = json.loads(f.read_text(encoding="utf-8")).get("generated_at")
            except (ValueError, OSError):
                pass
        entries.append({"key": loc["key"], "label": loc["label"],
                        "updated_at": updated})
    (data_dir / "index.json").write_text(
        json.dumps({"generated_at": now_iso(), "locations": entries}, indent=1),
        encoding="utf-8")


# ---------------------------------------------------------------- git

def git(repo_dir: str, *args, check=True):
    return subprocess.run(["git", *args], cwd=repo_dir, check=check,
                          capture_output=True, text=True)


def commit_and_push(repo_dir: str, message: str, push: bool = True) -> bool:
    git(repo_dir, "add", "data")
    if not git(repo_dir, "status", "--porcelain", "data").stdout.strip():
        log.info("No data changes to commit.")
        return False
    git(repo_dir, "commit", "-m", message)
    log.info("Committed: %s", message)
    if not push:
        return True
    pull = git(repo_dir, "pull", "--rebase", check=False)
    if pull.returncode != 0:
        log.warning("git pull --rebase failed:\n%s", pull.stderr.strip())
    try:
        git(repo_dir, "push")
        log.info("Pushed.")
        return True
    except subprocess.CalledProcessError as e:
        log.error("git push failed:\n%s", e.stderr.strip())
        return False


# ---------------------------------------------------------------- publish

def publish(cfg: dict, keys=None, push: bool = True) -> None:
    data_dir = Path(cfg["repo_dir"]) / "data"
    changed = []
    for loc in cfg["locations"]:
        if keys and loc["key"] not in keys:
            continue
        if not Path(loc["path"]).exists():
            log.error("Missing file for %s: %s", loc["label"], loc["path"])
            continue
        try:
            payload = build_payload(loc, cfg)
        except Exception as e:  # keep going for the other locations
            log.error("Failed to build %s: %s", loc["label"], e)
            continue
        if write_if_changed(data_dir / f"{loc['key']}.json", payload):
            changed.append(loc["label"])
            log.info("Updated %s (%d rows)", loc["label"], len(payload["rows"]))
        else:
            log.info("%s unchanged.", loc["label"])

    if changed:
        write_index(data_dir, cfg)
        stamp = datetime.now().strftime("%Y-%m-%d %H:%M")
        commit_and_push(cfg["repo_dir"],
                        f"Update {', '.join(changed)} ({stamp})", push=push)


# ---------------------------------------------------------------- watching

def watch(cfg: dict, push: bool = True) -> None:
    from watchdog.events import FileSystemEventHandler
    from watchdog.observers import Observer

    by_path = {norm_path(loc["path"]): loc["key"] for loc in cfg["locations"]}
    debounce = float(cfg.get("debounce_seconds", 5))
    pending: dict[str, threading.Timer] = {}
    work: "queue.Queue[str]" = queue.Queue()
    lock = threading.Lock()

    def schedule(key: str) -> None:
        with lock:
            t = pending.get(key)
            if t:
                t.cancel()
            t = threading.Timer(debounce, work.put, args=(key,))
            t.daemon = True
            pending[key] = t
            t.start()

    class Handler(FileSystemEventHandler):
        def _check(self, path):
            if not path:
                return
            key = by_path.get(norm_path(path))
            if key:
                log.info("Change detected: %s", Path(path).name)
                schedule(key)

        def on_modified(self, e):
            self._check(e.src_path)

        def on_created(self, e):
            self._check(e.src_path)

        def on_moved(self, e):  # Dropbox writes a temp file, then renames
            self._check(e.dest_path)

    def worker():
        while True:
            key = work.get()
            keys = {key}
            time.sleep(1)  # collect anything else that fired at the same time
            while True:
                try:
                    keys.add(work.get_nowait())
                except queue.Empty:
                    break
            try:
                publish(cfg, keys=keys, push=push)
            except Exception as e:
                log.exception("Publish failed: %s", e)

    threading.Thread(target=worker, daemon=True).start()

    observer = Observer()
    handler = Handler()
    for folder in sorted({str(Path(p).parent) for p in by_path}):
        observer.schedule(handler, folder, recursive=False)
        log.info("Watching %s", folder)
    observer.start()
    log.info("Publisher running. Save a workbook to publish. Ctrl+C to stop.")
    try:
        while True:
            time.sleep(1)
    except KeyboardInterrupt:
        pass
    finally:
        observer.stop()
        observer.join()


# ---------------------------------------------------------------- main

def main() -> int:
    ap = argparse.ArgumentParser(description="Budget-OS publisher")
    ap.add_argument("--once", action="store_true",
                    help="publish all locations once and exit")
    ap.add_argument("--no-push", action="store_true",
                    help="commit but do not push")
    args = ap.parse_args()

    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s %(levelname)s %(message)s",
        handlers=[logging.FileHandler(LOG_PATH, encoding="utf-8"),
                  logging.StreamHandler(sys.stdout)],
    )

    cfg = json.loads(CONFIG_PATH.read_text(encoding="utf-8"))
    repo = Path(cfg["repo_dir"])
    if not (repo / ".git").exists():
        log.error("repo_dir is not a git repo: %s", repo)
        return 1

    publish(cfg, push=not args.no_push)
    if not args.once:
        watch(cfg, push=not args.no_push)
    return 0


if __name__ == "__main__":
    sys.exit(main())
